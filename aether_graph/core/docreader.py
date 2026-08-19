"""Lectura determinista de documentos para la pasada de enriquecimiento.

Soporta PDF, DOCX y XLSX con imports OPCIONALES: si la librería no está
instalada (ej. MCP stdio en un contenedor mínimo), devuelve texto vacío sin
romper el flujo 100% stdlib de AetherGraph.
"""
from pathlib import Path

DOC_EXTS = (".pdf", ".docx", ".xlsx", ".xlsm")


def extract_document_text(path: Path, max_chars: int = 60000) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _pdf_text(path, max_chars)
        if suffix == ".docx":
            return _docx_text(path, max_chars)
        if suffix in (".xlsx", ".xlsm"):
            return _xlsx_text(path, max_chars)
    except Exception:
        return ""
    return ""


def _pdf_text(path: Path, max_chars: int) -> str:
    reader = None
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
    except Exception:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(str(path))
        except Exception:
            return ""
    out = []
    total = 0
    for page in reader.pages[:20]:
        t = page.extract_text() or ""
        out.append(t)
        total += len(t)
        if total >= max_chars:
            break
    return "\n".join(out)[:max_chars]


def _docx_text(path: Path, max_chars: int) -> str:
    try:
        from docx import Document
    except Exception:
        return ""
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)[:max_chars]


def _xlsx_text(path: Path, max_chars: int) -> str:
    try:
        import openpyxl
    except Exception:
        return ""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets[:5]:
        out.append(f"# Hoja: {ws.title}")
        for row in ws.iter_rows(max_row=30, values_only=True):
            cells = [str(c) for c in row if c is not None][:8]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)[:max_chars]
