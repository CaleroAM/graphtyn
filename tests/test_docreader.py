import pytest
from pathlib import Path

from aether_graph.core.ast_parser import ASTParser
from aether_graph.core.docreader import extract_document_text


def test_doc_reference_links(tmp_path):
    (tmp_path / "a.md").write_text("Ver [detalle](b.md) y [[otro]]\n", encoding="utf-8")
    (tmp_path / "b.md").write_text("Otro doc\n", encoding="utf-8")
    (tmp_path / "otro.md").write_text("wiki doc\n", encoding="utf-8")
    g = ASTParser().scan_directory(tmp_path, respect_git=False)
    refs = [(l["source"], l["target"]) for l in g["links"] if l.get("label") == "referencia"]
    assert ("file:a.md", "file:b.md") in refs
    assert ("file:a.md", "file:otro.md") in refs


def test_doc_reference_ignores_external(tmp_path):
    (tmp_path / "a.md").write_text("Enlace [web](https://example.com/x) y [img](img.png)\n", encoding="utf-8")
    g = ASTParser().scan_directory(tmp_path, respect_git=False)
    refs = [l for l in g["links"] if l.get("label") == "referencia"]
    assert refs == []


def test_pdf_docx_xlsx_nodes_created(tmp_path):
    (tmp_path / "manual.pdf").write_bytes(b"%PDF-1.4 basura")
    (tmp_path / "nota.docx").write_bytes(b"PK basura")
    (tmp_path / "datos.xlsx").write_bytes(b"PK basura")
    g = ASTParser().scan_directory(tmp_path, respect_git=False)
    files = {n["id"] for n in g["nodes"] if n["id"].startswith("file:")}
    assert "file:manual.pdf" in files
    assert "file:nota.docx" in files
    assert "file:datos.xlsx" in files


def test_docx_extraction(tmp_path):
    pytest.importorskip("docx")
    from docx import Document
    doc = Document()
    doc.add_paragraph("Resumen ejecutivo del proyecto 366metrics")
    p = tmp_path / "resumen.docx"
    doc.save(str(p))
    txt = extract_document_text(p)
    assert "Resumen ejecutivo" in txt


def test_xlsx_extraction(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "tenants"
    ws.append(["id", "nombre"])
    ws.append([1, "Eduardo"])
    p = tmp_path / "tenants.xlsx"
    wb.save(str(p))
    txt = extract_document_text(p)
    assert "tenants" in txt
    assert "Eduardo" in txt


def test_unreadable_document_returns_empty(tmp_path):
    p = tmp_path / "roto.pdf"
    p.write_bytes(b"esto no es un pdf real")
    assert extract_document_text(p) == ""


def test_image_nodes_get_image_kind(tmp_path):
    (tmp_path / "diagram.png").write_bytes(b"\x89PNG\r\n\x1a\nfalso")
    (tmp_path / "foto.jpg").write_bytes(b"\xff\xd8\xff\xe0falso")
    g = ASTParser().scan_directory(tmp_path, respect_git=False)
    kinds = {n["id"]: n["kind"] for n in g["nodes"] if n["id"].startswith("file:")}
    assert kinds.get("file:diagram.png") == "image"
    assert kinds.get("file:foto.jpg") == "image"


def test_media_nodes_get_media_kind(tmp_path):
    (tmp_path / "demo.mp4").write_bytes(b"\x00\x00\x00\x18ftypmp42falso")
    (tmp_path / "grabacion.mp3").write_bytes(b"ID3falso")
    g = ASTParser().scan_directory(tmp_path, respect_git=False)
    kinds = {n["id"]: n["kind"] for n in g["nodes"] if n["id"].startswith("file:")}
    assert kinds.get("file:demo.mp4") == "media"
    assert kinds.get("file:grabacion.mp3") == "media"


def test_transcribe_missing_lib_returns_empty(monkeypatch, tmp_path):
    import sys
    import types

    fake = types.ModuleType("faster_whisper")
    fake.WhisperModel = lambda *a, **k: (_ for _ in ()).throw(ImportError("no disponible"))
    monkeypatch.setitem(sys.modules, "faster_whisper", fake)
    p = tmp_path / "audio.mp3"
    p.write_bytes(b"basura")
    from aether_graph.core.docreader import transcribe_media
    assert transcribe_media(p) == ""
